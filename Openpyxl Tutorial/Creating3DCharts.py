#3D charts 

from openpyxl import load_workbook
from openpyxl.chart import Reference, Series, BarChart3D


wb = load_workbook(filename='chart.xlsx')
ws = wb.active

data = Reference(ws, min_col=2, min_row=1, max_col=3, max_row=4)
titles = Reference(ws, min_col=1, min_row=2, max_row=4)
chart = BarChart3D()
chart.title = "3D Bar Chart"
chart.add_data(data=data, titles_from_data=True)
chart.set_categories(titles)

ws.add_chart(chart, "E5")
wb.save("bar3d.xlsx")