#3dLin chart

from datetime import date

from openpyxl import load_workbook
from openpyxl.chart import LineChart3D, Reference

wb = load_workbook(filename='chart.xlsx')
ws = wb.active

c1 = LineChart3D()
c1.title = "3D Line Chart"
c1.legend = None
c1.style = 15
c1.y_axis.title = 'Size'
c1.x_axis.title = 'Test Number'

data = Reference(ws, min_col=2, min_row=1, max_col=4, max_row=7)
c1.add_data(data, titles_from_data=True)

ws.add_chart(c1, "A10")

wb.save("line3D.xlsx")