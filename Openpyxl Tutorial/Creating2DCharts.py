# creating  2 D charts from an existing file
# see creating bar charts to creat a benad new file and hashtag comments at top

from openpyxl import load_workbook
from openpyxl.chart import AreaChart, Reference

wb = load_workbook(filename='chart.xlsx')
ws = wb.active

chart = AreaChart()
chart.title = "Defects per Aircraft"
chart.style = 10
chart.x_axis.title = 'Weeks'
chart.y_axis.title = 'Defect Count'

cats = Reference(ws, min_col=1, min_row=1, max_row=7)
data = Reference(ws, min_col=2, min_row=1, max_col=3, max_row=7)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)

ws.add_chart(chart, "A5")

wb.save("AreaChart.xlsx")