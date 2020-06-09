#line Charts in a new sheet

from openpyxl import  load_workbook
from openpyxl.chart import LineChart, Reference, Series


wb = load_workbook(filename='chart.xlsx')
ws = wb.active
cs = wb.create_chartsheet()

chart = LineChart()
labels = Reference(ws, min_col=1, min_row=1, max_row=3)
data = Reference(ws, min_col=2, min_row=1, max_row=3)
chart.series = (Series(data),)
chart.title = "LineChart"

cs.add_chart(chart)

wb.save("Sheetdemo.xlsx")